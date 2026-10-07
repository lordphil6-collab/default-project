"""Upload parsing, notify text, paystack signature — pure unit tests (no servers)."""
import hashlib
import hmac
import io

import pytest

from backend.app.services import notify as n
from backend.app.services import parse_files as pf
from backend.app.routers import billing as b


def test_extract_text_plain():
    assert pf.extract_text("q.txt", "text/plain", b"Ocean Freight $100") == "Ocean Freight $100"


def test_extract_excel_and_word():
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["Ocean Freight", 1850])
    buf = io.BytesIO()
    wb.save(buf)
    assert "1850" in pf.extract_text("q.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", buf.getvalue())

    import docx

    doc = docx.Document()
    doc.add_paragraph("Destination Charges $650")
    buf2 = io.BytesIO()
    doc.save(buf2)
    assert "650" in pf.extract_text("q.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document", buf2.getvalue())


def test_extract_pdf():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Transit 35 days")
    assert "35 days" in pf.extract_text("q.pdf", "application/pdf", doc.tobytes())


def test_rejects_images_and_empty():
    with pytest.raises(ValueError, match="Tesseract"):
        pf.extract_text("q.png", "image/png", b"12345")
    with pytest.raises(ValueError, match="empty"):
        pf.extract_text("q.txt", "text/plain", b"")


def test_quote_text_and_logged_only_delivery():
    body = n.quote_text("Acme", "China>Lagos", 3270.4, "USD", 7, "")
    assert "3,270.40" in body
    assert n.send_email("a@b.c", "s", "b")["status"] == "logged-only"
    assert n.send_whatsapp("+234", "b")["status"] == "logged-only"


def test_paystack_signature():
    secret, payload = "s3cret", b'{"event":"charge.success"}'
    sig = hmac.new(secret.encode(), payload, hashlib.sha512).hexdigest()
    assert b.verify_paystack_signature(secret, payload, sig) is True
    assert b.verify_paystack_signature(secret, payload, "bad") is False
    assert b.verify_paystack_signature("other", payload, sig) is False
