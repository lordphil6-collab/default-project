"""File text extraction for quotation ingest — OSS only (free, no cloud OCR).

Supported: PDF (pymupdf), Excel (openpyxl, all sheets), Word (python-docx),
plain text/CSV. Images need a Tesseract binary (not installed) and are
rejected with a clear 422 telling the user to paste the text instead.
Max 10 MB per file.
"""
import io

MAX_BYTES = 10 * 1024 * 1024
IMAGE_TYPES = {"image/png", "image/jpeg", "image/jpg", "image/webp", "image/tiff"}


def extract_text(filename: str, content_type: str, data: bytes) -> str:
    if len(data) > MAX_BYTES:
        raise ValueError("file over 10 MB — split it or paste the text")
    if not data:
        raise ValueError("empty file")
    name = (filename or "").lower()
    if content_type in IMAGE_TYPES or name.endswith((".png", ".jpg", ".jpeg", ".webp", ".tiff", ".bmp")):
        raise ValueError("image OCR needs a Tesseract binary (not installed) — paste the text instead")
    if name.endswith(".pdf") or content_type == "application/pdf":
        import pymupdf

        doc = pymupdf.open(stream=data, filetype="pdf")
        return "\n".join(page.get_text() for page in doc)
    if name.endswith((".xlsx", ".xlsm", ".xltx")) or content_type in (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                out.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    if name.endswith((".docx",)) or content_type in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",):
        import docx

        doc = docx.Document(io.BytesIO(data))
        paras = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                paras.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paras)
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValueError("unrecognized format — send PDF, Excel, Word, or plain text")
